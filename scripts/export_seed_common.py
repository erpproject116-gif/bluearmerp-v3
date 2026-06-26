"""Shared helpers for Excel export → demo SQL seed generators."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path

import openpyxl


def sql_str(s: str | None) -> str:
    if s is None:
        return "null"
    return "'" + str(s).replace("'", "''") + "'"


def parse_date_no(raw: str) -> tuple[date, int]:
    raw = str(raw).strip()
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})\s*-\s*(\d+)$", raw)
    if not m:
        raise ValueError(f"bad date-no: {raw!r}")
    month, day, year, seq = m.groups()
    return date(int(year), int(month), int(day)), int(seq)


def parse_us_date(raw: str | None) -> date | None:
    if not raw:
        return None
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", str(raw).strip())
    if not m:
        return None
    month, day, year = m.groups()
    return date(int(year), int(month), int(day))


def parse_money(raw) -> float | None:
    if raw is None or raw == "":
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    s = str(raw).strip().replace(",", "")
    if not s:
        return None
    return float(s)


def normalize_item_code(raw: str, seen: set[str]) -> str:
    """Map legacy export codes to inv_items char(5)."""
    s = str(raw).strip().upper()
    if not s:
        raise ValueError("empty item code")
    if len(s) <= 5:
        code = s.zfill(5) if s.isdigit() else s.ljust(5)[:5]
    else:
        code = s[-5:]
    if code in seen:
        n = 80001
        while str(n)[-5:] in seen or f"{n:05d}" in seen:
            n += 1
        code = f"{n:05d}"[-5:] if n < 100000 else f"X{n % 10000:04d}"[:5]
    seen.add(code)
    return code


def find_sheet(wb: openpyxl.Workbook, preferred: list[str], header_needle: str) -> str:
    lowered = {name.lower(): name for name in wb.sheetnames}
    for want in preferred:
        if want.lower() in lowered:
            return lowered[want.lower()]
    for name in wb.sheetnames:
        ws = wb[name]
        for row in ws.iter_rows(min_row=1, max_row=5, values_only=True):
            flat = " ".join(str(c) for c in row if c)
            if header_needle.lower() in flat.lower():
                return name
    raise KeyError(f"no sheet matching {preferred!r} / {header_needle!r} in {wb.sheetnames}")


def header_row_index(ws, needles: tuple[str, ...]) -> int:
    for idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        flat = " ".join(str(c) for c in row if c).lower()
        if all(n.lower() in flat for n in needles):
            return idx
    return 2


def col_index(headers: tuple, *names: str) -> int | None:
    norm = [str(h or "").strip().lower() for h in headers]
    for name in names:
        key = name.lower()
        for i, h in enumerate(norm):
            if key in h or h in key:
                return i
    return None
